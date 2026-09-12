"""Shared section order for entry, preview and document exports."""
from io import BytesIO
from xml.sax.saxutils import escape

from app.i18n import translate_text, translate_value
from app.services.department_reports import clean_lines, report_period_text


SECTIONS = {
    "maintenance": [
        ("team", "Equipa em serviço", []),
        ("equipment_status", "Equipamentos inspecionados", []),
        ("activities", "Trabalhos executados", []),
        ("incidents", "Paragens e emergências", []),
        ("readings", "Utilidades e equipamentos críticos", ["Bomba do furo 1", "Bomba do furo 2", "Bomba de incêndio 1", "Bomba de incêndio 2"]),
        ("notes", "Notas adicionais", []),
        ("pending_actions", "Pendências / ações para acompanhamento", []),
    ],
    "security": [
        ("incidents", "Incidentes", []),
        ("equipment_status", "Apreensões e meios operacionais", ["Entrada de mercadorias apreendidas", "Entrada de viaturas apreendidas", "Saída de viaturas apreendidas", "Saída de mercadorias apreendidas", "Meios operacionais"]),
        ("notes", "Outras informações", []),
        ("team", "Staff da segurança", ["Faltas GT", "Faltas G4S", "Férias GT", "Presenças na parada conjunta"]),
        ("activities", "Distribuição de postos e patrulhas", ["Equipa GT", "Equipa G4S", "Patrulha diurna", "Patrulha nocturna"]),
        ("readings", "Iluminação e leituras", ["Iluminação e cortes de energia", "Geradores", "EDM", "Centro Social", "Mects", "Bypass", "Bombas"]),
        ("pending_actions", "Pendências / ações para acompanhamento", []),
    ],
    "it": [
        ("team", "Equipa em serviço", []),
        ("activities", "Atividades dos ITs / suporte executado", []),
        ("incidents", "Incidentes técnicos / falhas de sistemas", []),
        ("equipment_status", "Equipamentos, redes e sistemas verificados", []),
        ("readings", "Indicadores técnicos / tickets / disponibilidade", []),
        ("pending_actions", "Pendências / ações para acompanhamento", []),
        ("notes", "Notas adicionais", []),
    ],
}


def collect_sections(form, key, values):
    # Store readable labelled text in existing columns; legacy reports remain intact.
    for field, _, subtitles in SECTIONS[key]:
        blocks = []
        for index, subtitle in enumerate(subtitles):
            value = str(form.get(f"{field}__{index}", "")).strip()
            if value:
                blocks.append(f"{subtitle}:\n{value}")
        if blocks:
            values[field] = "\n\n".join([str(values.get(field) or "").strip(), *blocks]).strip()
    return values


def presentation(report, language="pt"):
    from app.services.department_tables import present_tables
    tables = present_tables(report)
    names = {"maintenance": "Departamento de Manutenção", "security": "Departamento de Proteção e Segurança", "it": "Departamento de Informática"}
    tx = lambda value: translate_text(value, language)
    return {
        "id": report.id, "number": report.number,
        "title": tx(names[report.department_key]),
        "meta": [(tx(label), value or "—") for label, value in [
            ("Data", report.report_date.strftime("%d/%m/%Y")),
            ("Turno / período", report_period_text(report) if report.shift or report.period_start or report.period_end else tx("Não informado")),
            ("Preparado por", report.prepared_by or report.created_by.full_name),
            ("Supervisor", report.supervisor), ("Local / área", report.location),
            ("Estado", translate_value(report.status, language)),
        ]],
        "sections": [{"title": tx(title), "lines": [
            {"text": tx(line[:-1]) if line.endswith(":") else line, "heading": line.endswith(":")}
            for line in clean_lines(getattr(report, field))
        ] or ([] if any(t['section'] == field for t in tables) else [{"text": tx("Não informado"), "heading": False}]), 'tables': [t for t in tables if t['section'] == field]} for field, title, _ in SECTIONS[report.department_key]],
    }


def reports_pdf(items, title, generated_by, language="pt"):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Image, KeepTogether
    from app.services.pdf_branding import brand_header, branded_styles, branded_footer, GOLD_DARK
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=1.8*cm, rightMargin=1.8*cm, topMargin=1.5*cm, bottomMargin=1.6*cm)
    styles, _, _ = branded_styles()
    styles["Normal"].fontSize = 10
    styles["Normal"].leading = 15
    styles["GTSection"].borderColor = GOLD_DARK
    styles["GTSection"].borderWidth = 0.5
    styles["GTSection"].borderPadding = 6
    styles["GTSection"].spaceAfter = 12
    styles["GTSection"].keepWithNext = True
    styles["Heading3"].spaceBefore = 8
    styles["Heading3"].spaceAfter = 4
    story = []
    for item in items:
        if story: story.append(PageBreak())
        story.append(brand_header(escape(item["title"]), subtitle=escape(title) if title != item['title'] else '', meta=[escape(item["number"])], width=doc.width))
        story.append(Spacer(1, 12))
        for label, value in item["meta"]:
            story.append(Paragraph(f'<b>{escape(label)}</b>: {escape(value)}', styles["Normal"]))
        for index, section in enumerate(item["sections"], 1):
            story.append(Paragraph(f'{index:02d}  {escape(section["title"])}', styles["GTSection"]))
            for line in section["lines"]:
                style = styles["Heading3"] if line["heading"] else styles["Normal"]
                story.append(Paragraph(escape(line["text"]), style))
                story.append(Spacer(1, 4))
            for table in section.get('tables', []):
                from reportlab.platypus import Table, TableStyle
                from reportlab.lib import colors
                from reportlab.lib.styles import ParagraphStyle
                cell_style = ParagraphStyle('TableCell', parent=styles['Normal'], fontSize=8, leading=11)
                head_style = ParagraphStyle('TableHeader', parent=cell_style, fontName='Helvetica-Bold')
                story.append(Paragraph(escape(table['title']), styles['Heading3']))
                if table.get('note'): story.append(Paragraph(escape(table['note']), styles['Normal']))
                if table['rows']:
                    cells = [[Paragraph(escape(str(c)).replace('\n', '<br/>'), head_style if ri == 0 else cell_style) for c in row] for ri, row in enumerate([table['columns'], *table['rows']])]
                    weights = [.6,2.2,1,1.2,1.6,1.5,1.3] if table.get('key') == 'parade' and len(table['columns']) == 7 else [1]*len(table['columns'])
                    grid = Table(cells, colWidths=[doc.width*w/sum(weights) for w in weights], repeatRows=1, hAlign='LEFT', splitInRow=int(any(len(c)>500 for row in table['rows'] for c in row)))
                    grid.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.HexColor('#FFF4CF')), ('GRID', (0,0), (-1,-1), .4, colors.HexColor('#B8BABD')), ('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 5), ('RIGHTPADDING', (0,0), (-1,-1), 5), ('TOPPADDING', (0,0), (-1,-1), 5), ('BOTTOMPADDING', (0,0), (-1,-1), 5)]))
                    story.extend([grid, Spacer(1, 8)])
            if section['title'] == 'Indicadores operacionais':
                from app.services.reporting_charts import metric_chart
                for metric in item.get('charts', [])[:6]:
                    story.append(KeepTogether([Paragraph(escape(metric['label']+' / '+metric['dimension']+' ('+metric['unit']+')'), styles['Heading3']), Image(BytesIO(metric_chart(metric)), width=doc.width, height=doc.width*340/1100), Paragraph('Dias sem observação não são representados como zero.', styles['Normal'])]))
    if not story: story.append(Paragraph(translate_text("Sem dados para apresentar.", language), styles["Normal"]))
    footer = lambda c, d: branded_footer(c, d, generated_by, language)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def reports_docx(items, title, generated_by, language="pt"):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from app.config import get_settings
    doc = Document()
    # Strip template theme rules so the document uses only GT branding.
    for node in list(doc.styles.element.iter(qn("w:pBdr"))):
        node.getparent().remove(node)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.6), Cm(1.7)
    section.header_distance = section.footer_distance = Cm(.8)
    section.left_margin = section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.color.rgb = RGBColor.from_string("2D3033")
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Heading 2"].font.size = Pt(11)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(6)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(4)
    doc.styles['Heading 2'].paragraph_format.keep_with_next = True
    logo = get_settings().logo_path
    if logo.exists(): section.header.paragraphs[0].add_run().add_picture(str(logo), width=Cm(3))
    footer = section.footer.paragraphs[0]
    footer.add_run("Gestão de Terminais, SA · "+translate_text("Página", language)+" ")
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); footer._p.append(fld)
    for index, item in enumerate(items):
        if index: doc.add_page_break()
        doc.add_paragraph(item["title"], "Title")
        doc.add_paragraph((title+" · " if title != item['title'] else '')+item["number"])
        for label, value in item["meta"]:
            p = doc.add_paragraph(); p.add_run(label+": ").bold = True; p.add_run(value)
        for number, block in enumerate(item["sections"], 1):
            p = doc.add_paragraph(f'{number:02d}  {block["title"]}', "Heading 1")
            shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "FFF4CF"); p._p.get_or_add_pPr().append(shading)
            for line in block["lines"]:
                doc.add_paragraph(line["text"], "Heading 2" if line["heading"] else "Normal")
            for table in block.get('tables', []):
                doc.add_paragraph(table['title'], 'Heading 2')
                if table.get('note'): doc.add_paragraph(table['note'])
                if table['rows']:
                    grid = doc.add_table(rows=1, cols=len(table['columns']))
                    grid.style = 'Table Grid'
                    repeat = OxmlElement('w:tblHeader'); grid.rows[0]._tr.get_or_add_trPr().append(repeat)
                    for cell, text in zip(grid.rows[0].cells, table['columns']):
                        cell.text = text
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'FFF4CF'); cell._tc.get_or_add_tcPr().append(shade)
                    for row in table['rows']:
                        for cell, text in zip(grid.add_row().cells, row): cell.text = text
                    for ri, row in enumerate(grid.rows):
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                p.paragraph_format.space_after = Pt(3)
                                for run in p.runs: run.font.size = Pt(8); run.bold = ri == 0
                    doc.add_paragraph()
            if block['title'] == 'Indicadores operacionais':
                from app.services.reporting_charts import metric_chart
                for metric in item.get('charts', [])[:6]:
                    doc.add_paragraph(metric['label']+' / '+metric['dimension']+' ('+metric['unit']+')', 'Heading 2')
                    doc.add_picture(BytesIO(metric_chart(metric)), width=Cm(17))
                    doc.add_paragraph('Dias sem observação não são representados como zero.')
    if not items: doc.add_paragraph(translate_text("Sem dados para apresentar.", language))
    doc.core_properties.author = generated_by
    doc.core_properties.title = title
    out = BytesIO(); doc.save(out); return out.getvalue()
